"""CSV/XLSX data uploads for cross-verification against API-sourced emissions."""

import logging
from datetime import date, datetime
from uuid import UUID

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from sqlalchemy.orm import Session

from database import get_db
from dependencies import check_facility_access, require_role
from models.facility import Facility
from models.upload import Upload
from models.user import User
from schemas.upload import UploadPreviewResponse, UploadResponse
from services.uploads import (
    convert_xlsx_to_csv_bytes,
    parse_emissions_csv,
    stage_upload_drafts,
)

logger = logging.getLogger(__name__)

router = APIRouter()

PREVIEW_COLUMNS = ["timestamp", "metric", "value", "unit", "facility_name"]
PREVIEW_ROW_LIMIT = 5
# Curated set of columns worth previewing from an electricity workbook.
# Anything not present in the sheet is simply skipped.
PREVIEW_IMPORTANT_HEADERS = {
    "month",
    "total units",
    "net units",
    "net bill",
    "per unit",
    "unit used",
}


def _display_cell(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        rounded = round(value, 2)
        return str(int(rounded)) if rounded.is_integer() else str(rounded)
    text = str(value).strip()
    return text or None


def preview_workbook(
    content: bytes, facility_name: str | None = None
) -> tuple[list[str], list[list[str | None]], int]:
    """Curated preview of an electricity workbook: important columns + rows.

    Returns ``(columns, rows, row_count)``. ``columns`` is empty when the
    sheet has none of the important headers, signalling the caller to fall
    back to the canonical CSV preview.
    """
    from io import BytesIO

    from openpyxl import load_workbook

    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.active
    raw = [list(row) for row in sheet.iter_rows(values_only=True)]
    workbook.close()

    if not raw:
        return [], [], 0

    headers = [
        (index, str(cell).strip().lower() if cell is not None else "")
        for index, cell in enumerate(raw[0])
    ]
    selected = [index for index, header in headers if header in PREVIEW_IMPORTANT_HEADERS]
    if not selected:
        return [], [], 0

    columns = [
        str(raw[0][index]).strip() if raw[0][index] is not None else f"Column {index + 1}"
        for index in selected
    ]
    data_rows = [
        row
        for row in raw[1:]
        if any(cell is not None and str(cell).strip() != "" for cell in row)
    ]
    columns.append("Facility")
    rows = [
        [_display_cell(row[index]) for index in selected] + [facility_name]
        for row in data_rows[:PREVIEW_ROW_LIMIT]
    ]
    return columns, rows, len(data_rows)


@router.post("/preview", response_model=UploadPreviewResponse)
async def preview_upload(
    file: UploadFile = File(...),
    facility_id: UUID | None = Form(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_role("admin", "facility_manager")),
) -> UploadPreviewResponse:
    """Parse a CSV/XLSX and return its columns plus the first few rows.

    XLSX files show the real workbook's important columns (month/units/billing)
    plus the facility the rows will be assigned to. Runs the same validation as
    the real ingest, so it doubles as a dry run: invalid files get a 422 with
    the parse error instead of being ingested.
    """
    content = await file.read()
    is_xlsx = (file.filename or "").lower().endswith(".xlsx")

    try:
        csv_bytes = convert_xlsx_to_csv_bytes(content) if is_xlsx else content
        rows = parse_emissions_csv(csv_bytes)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)
        ) from exc
    except Exception as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail=f"Failed to parse xlsx: {exc}",
        ) from exc

    if is_xlsx:
        target_facility_id = facility_id or user.facility_id
        facility = db.get(Facility, target_facility_id) if target_facility_id else None
        try:
            columns, preview_rows, row_count = preview_workbook(
                content, facility.name if facility else None
            )
        except Exception as exc:
            # A malformed-but-parseable workbook must not surface as a 500 with
            # no body; fall back to the canonical CSV preview instead.
            logger.warning("workbook preview failed for %s: %s", file.filename, exc)
            columns, preview_rows, row_count = [], [], 0
        if columns:
            return UploadPreviewResponse(
                filename=file.filename or "upload.xlsx",
                columns=columns,
                rows=preview_rows,
                row_count=row_count,
            )

    return UploadPreviewResponse(
        filename=file.filename or "upload.csv",
        columns=PREVIEW_COLUMNS,
        rows=[
            [_display_cell(row.get(column)) for column in PREVIEW_COLUMNS]
            for row in rows[:PREVIEW_ROW_LIMIT]
        ],
        row_count=len(rows),
    )


@router.post("", response_model=UploadResponse, status_code=status.HTTP_201_CREATED)
async def upload_csv(
    file: UploadFile = File(...),
    facility_id: UUID | None = Form(None),
    db: Session = Depends(get_db),
    user: User = Depends(require_role("admin", "facility_manager")),
) -> Upload:
    """Stage a CSV/XLSX for review.

    Rows become ``OCRDraft`` records (source_type="upload") that a reviewer
    confirms on the intake page — the same lifecycle OCR and Excel intake use.
    Nothing reaches ``activity_records``/``calculated_emissions`` until then, so
    an unreviewed upload can't move a dashboard. Rows already staged or
    confirmed are skipped, so re-uploading a file is a no-op rather than a
    double count.

    ``facility_id`` is the facility picked in the intake UI; it defaults to the
    facility attached to the account.
    """
    if facility_id is not None:
        check_facility_access(user, facility_id)
    target_facility = facility_id or user.facility_id
    if target_facility is None:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
            detail="no facility assigned to your account — ask an admin to set one",
        )

    content = await file.read()

    # detect xlsx and convert to CSV for pipeline
    if (file.filename or "").lower().endswith(".xlsx"):
        try:
            content = convert_xlsx_to_csv_bytes(content)
        except Exception as exc:
            raise HTTPException(
                status_code=status.HTTP_422_UNPROCESSABLE_ENTITY,
                detail=f"Failed to parse xlsx: {exc}",
            ) from exc

    try:
        rows = parse_emissions_csv(content)
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)
        ) from exc

    try:
        drafts = stage_upload_drafts(
            db,
            user_id=user.id,
            facility_id=target_facility,
            filename=file.filename or "upload.csv",
            rows=rows,
        )
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=str(exc)
        ) from exc

    record = Upload(
        user_id=user.id,
        facility_id=target_facility,
        filename=file.filename or "upload.csv",
        status="pending_review",
        row_count=len(drafts),
    )
    db.add(record)
    db.commit()
    db.refresh(record)
    return record
